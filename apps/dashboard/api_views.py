from datetime import datetime
from io import BytesIO

from django.http import HttpResponse
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework.permissions import IsAuthenticated
from rest_framework import generics, status
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.models import User
from apps.accounts.serializers import UserSerializer
from apps.dataset.models import DatasetImage
from apps.dataset.serializers import DatasetImageAdminUpdateSerializer, DatasetImageSerializer
from apps.diagnosis.models import DetectionRecord, ModelConfig
from apps.diagnosis.serializers import DetectionRecordSerializer, ModelConfigSerializer

from .permissions import IsAdminRole
from .services import build_detection_statistics, build_filtered_detection_records


TITLE_FILL = PatternFill("solid", fgColor="DCEBD7")
HEADER_FILL = PatternFill("solid", fgColor="EDF5E8")
TITLE_FONT = Font(bold=True, size=12)
HEADER_FONT = Font(bold=True)


def _autosize_worksheet(worksheet):
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 4, 12), 40)


def _append_section_title(worksheet, title):
    row = 1 if worksheet.max_row == 1 and worksheet["A1"].value is None else worksheet.max_row + 1
    worksheet.cell(row=row, column=1, value=title)
    worksheet.cell(row=row, column=1).font = TITLE_FONT
    worksheet.cell(row=row, column=1).fill = TITLE_FILL


def _append_header(worksheet, headers):
    worksheet.append(headers)
    row = worksheet.max_row
    for index in range(1, len(headers) + 1):
        cell = worksheet.cell(row=row, column=index)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _append_table(worksheet, title, headers, rows):
    if worksheet.max_row > 1 or worksheet["A1"].value is not None:
        worksheet.append([])
    _append_section_title(worksheet, title)
    _append_header(worksheet, headers)
    for row in rows:
        worksheet.append(row)


class AdminDetectionListView(generics.ListAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = DetectionRecordSerializer
    queryset = DetectionRecord.objects.select_related("user", "advice").prefetch_related("results")


class AdminDatasetListView(generics.ListAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = DatasetImageSerializer
    queryset = DatasetImage.objects.select_related("user", "source_detection").filter(source_detection__isnull=False)


class AdminDatasetUpdateView(generics.UpdateAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = DatasetImageAdminUpdateSerializer
    queryset = DatasetImage.objects.filter(source_detection__isnull=False)


class AdminModelConfigListView(generics.ListAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = ModelConfigSerializer
    queryset = ModelConfig.objects.all()


class AdminModelActivateView(APIView):
    permission_classes = [IsAdminRole]

    def patch(self, request, pk):
        config = generics.get_object_or_404(ModelConfig, pk=pk)
        ModelConfig.objects.update(is_active=False)
        config.is_active = True
        config.save(update_fields=["is_active", "updated_at"])
        return Response(ModelConfigSerializer(config).data)


class AdminUserListView(generics.ListAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = UserSerializer
    queryset = User.objects.all().order_by("-date_joined")


class StatisticsRequestMixin:
    permission_classes = [IsAuthenticated]

    def _parse_filters(self, request):
        start_date = parse_date(request.query_params.get("start_date", ""))
        end_date = parse_date(request.query_params.get("end_date", ""))
        if request.query_params.get("start_date") and start_date is None:
            return None, Response({"detail": "开始日期格式无效。"}, status=status.HTTP_400_BAD_REQUEST)
        if request.query_params.get("end_date") and end_date is None:
            return None, Response({"detail": "结束日期格式无效。"}, status=status.HTTP_400_BAD_REQUEST)
        if start_date and end_date and start_date > end_date:
            return None, Response({"detail": "开始日期不能晚于结束日期。"}, status=status.HTTP_400_BAD_REQUEST)

        user_id = request.query_params.get("user_id")
        filters = {
            "viewer": request.user,
            "start_date": start_date,
            "end_date": end_date,
            "category": request.query_params.get("category", "").strip(),
            "status": request.query_params.get("status", "").strip(),
            "model_key": request.query_params.get("model", "").strip(),
            "user_id": int(user_id) if user_id and user_id.isdigit() else None,
            "granularity": request.query_params.get("granularity", "auto").strip() or "auto",
            "include_recent": request.query_params.get("include_recent", "1").strip().lower()
            not in {"0", "false", "no"},
        }
        return filters, None


class StatisticsView(StatisticsRequestMixin, APIView):
    def get(self, request):
        filters, error_response = self._parse_filters(request)
        if error_response is not None:
            return error_response
        payload = build_detection_statistics(
            **{**filters, "include_recent": False, "fast": True}
        )
        return Response(payload)


class AdminStatisticsView(StatisticsView):
    permission_classes = [IsAdminRole]


class StatisticsExportView(StatisticsRequestMixin, APIView):
    def get(self, request):
        filters, error_response = self._parse_filters(request)
        if error_response is not None:
            return error_response

        payload = build_detection_statistics(**filters)
        records = build_filtered_detection_records(
            viewer=request.user,
            start_date=filters["start_date"],
            end_date=filters["end_date"],
            category=filters["category"],
            status=filters["status"],
            model_key=filters["model_key"],
            user_id=filters["user_id"],
        ).prefetch_related("results")

        workbook = Workbook()
        overview_sheet = workbook.active
        overview_sheet.title = "概览"
        overview_sheet.freeze_panes = "A2"
        overview_sheet["A1"] = "检测统计报表"
        overview_sheet["A1"].font = Font(bold=True, size=14)
        overview_sheet["A2"] = "导出时间"
        overview_sheet["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        _append_table(
            overview_sheet,
            "筛选条件",
            ["条件", "值"],
            [
                ["开始日期", payload["filters"]["selected"]["start_date"] or "全部"],
                ["结束日期", payload["filters"]["selected"]["end_date"] or "全部"],
                ["病虫害类别", payload["filters"]["selected"]["category"] or "全部"],
                ["检测状态", payload["filters"]["selected"]["status"] or "全部"],
                ["模型版本", payload["filters"]["selected"]["model"] or "全部"],
                ["用户 ID", payload["filters"]["selected"]["user_id"] or "全部"],
                ["时间粒度", payload["filters"]["selected"]["granularity"] or "auto"],
            ],
        )
        _append_table(
            overview_sheet,
            "摘要指标",
            ["指标", "数值"],
            [[key, value] for key, value in payload["summary"].items()],
        )

        timeline_sheet = workbook.create_sheet("趋势")
        _append_table(
            timeline_sheet,
            "时间趋势",
            ["时间", "总记录数", "成功数", "失败数"],
            [[item["label"], item["total"], item["success"], item["failed"]] for item in payload["charts"]["timeline"]],
        )

        category_sheet = workbook.create_sheet("类别分布")
        _append_table(
            category_sheet,
            "类别分布",
            ["类别", "次数", "平均置信度(%)"],
            [[item["label"], item["total"], item["avg_confidence"]] for item in payload["charts"]["categories"]],
        )

        status_sheet = workbook.create_sheet("状态分布")
        _append_table(
            status_sheet,
            "状态分布",
            ["状态", "数量", "占比(%)"],
            [[item["label"], item["total"], item["percent"]] for item in payload["charts"]["status_distribution"]],
        )

        model_sheet = workbook.create_sheet("模型分布")
        _append_table(
            model_sheet,
            "模型分布",
            ["模型", "记录数"],
            [[item["label"], item["total"]] for item in payload["charts"]["model_distribution"]],
        )

        user_sheet = workbook.create_sheet("用户排行")
        _append_table(
            user_sheet,
            "用户排行",
            ["用户", "记录数", "成功数"],
            [[item["label"], item["total"], item["success"]] for item in payload["charts"]["user_distribution"]],
        )

        confidence_sheet = workbook.create_sheet("置信度分布")
        _append_table(
            confidence_sheet,
            "置信度分布",
            ["区间", "数量"],
            [[item["label"], item["total"]] for item in payload["charts"]["confidence_distribution"]],
        )

        detail_sheet = workbook.create_sheet("检测明细")
        _append_table(
            detail_sheet,
            "检测记录明细",
            ["记录ID", "用户", "状态", "模型", "创建时间", "耗时(ms)", "原图路径", "检测框数量", "检测类别"],
            [
                [
                    record.id,
                    record.user.username,
                    record.get_status_display(),
                    f"{record.model_name or '未记录'} / {record.model_version or '-'}",
                    record.created_at.strftime("%Y-%m-%d %H:%M:%S"),
                    record.duration_ms,
                    f"detection-{record.id}-original.jpg",
                    record.results.count(),
                    " | ".join(sorted({result.label for result in record.results.all()})) or "-",
                ]
                for record in records.order_by("-created_at")
            ],
        )

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = worksheet.freeze_panes or "A2"
            _autosize_worksheet(worksheet)

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        filename = f'detection_statistics_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class AdminStatisticsExportView(StatisticsExportView):
    permission_classes = [IsAdminRole]
