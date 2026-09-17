import base64
from io import BytesIO
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.db import connection
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from openpyxl import load_workbook
from PIL import Image
from rest_framework.test import APIClient

from apps.diagnosis.models import DetectionRecord, DetectionResult

User = get_user_model()


def build_test_image(name="leaf.jpg", color="green"):
    buffer = BytesIO()
    image = Image.new("RGB", (200, 160), color=color)
    image.save(buffer, format="JPEG")
    return SimpleUploadedFile(name, buffer.getvalue(), content_type="image/jpeg")


def build_test_image_base64(color="green"):
    buffer = BytesIO()
    image = Image.new("RGB", (200, 160), color=color)
    image.save(buffer, format="JPEG")
    return base64.b64encode(buffer.getvalue()).decode("ascii")


@override_settings(MEDIA_ROOT="test_media")
class AdminApiTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_user(username="admin_user", password="secret123", role="admin")
        self.client = APIClient()

    def test_non_admin_cannot_access_admin_api(self):
        user = User.objects.create_user(username="normal")
        self.client.force_authenticate(user=user)
        response = self.client.get("/api/admin/users/")
        self.assertEqual(response.status_code, 403)

    def test_admin_can_access_dashboard_page(self):
        self.client.logout()
        browser = self.client
        browser.login(username="admin_user", password="secret123")
        response = browser.get("/dashboard/")
        self.assertEqual(response.status_code, 200)

    def test_dashboard_login_redirects_to_dashboard_index(self):
        response = self.client.post(
            "/dashboard/login/",
            {"username": "admin_user", "password": "secret123"},
        )
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response.url, "/dashboard/")

    def test_root_homepage_is_available(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)

    def test_admin_can_access_statistics_page(self):
        self.client.logout()
        browser = self.client
        browser.login(username="admin_user", password="secret123")
        response = browser.get("/dashboard/statistics/")
        self.assertEqual(response.status_code, 200)

    def test_admin_statistics_api_supports_filters(self):
        user = User.objects.create_user(username="farmer01")
        old_record = DetectionRecord.objects.create(
            user=user,
            original_image=build_test_image_base64(),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=120,
        )
        old_record.created_at = timezone.now() - timedelta(days=10)
        old_record.save(update_fields=["created_at"])
        DetectionResult.objects.create(
            record=old_record,
            label="leaf_blast",
            confidence=0.91,
            xmin=1,
            ymin=1,
            xmax=10,
            ymax=10,
        )

        new_record = DetectionRecord.objects.create(
            user=user,
            original_image=build_test_image_base64(color="red"),
            status="failed",
            model_name="rice-pest-detector",
            model_version="v2",
            duration_ms=180,
        )
        new_record.created_at = timezone.now()
        new_record.save(update_fields=["created_at"])
        DetectionResult.objects.create(
            record=new_record,
            label="brown_spot",
            confidence=0.62,
            xmin=2,
            ymin=2,
            xmax=12,
            ymax=12,
        )

        self.client.force_authenticate(user=self.admin)
        today = timezone.localdate().isoformat()
        response = self.client.get(f"/api/admin/statistics/?start_date={today}&category=brown_spot")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["summary"]["total_records"], 1)
        self.assertEqual(response.data["summary"]["failed_records"], 1)
        self.assertEqual(response.data["charts"]["categories"][0]["label"], "brown_spot")
        self.assertEqual(response.data["charts"]["model_distribution"][0]["label"], "rice-pest-detector / v2")

    def test_authenticated_user_statistics_api_only_returns_own_data(self):
        own_user = User.objects.create_user(username="owner01")
        other_user = User.objects.create_user(username="other01")
        DetectionRecord.objects.create(
            user=own_user,
            original_image=build_test_image_base64(),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=90,
        )
        DetectionRecord.objects.create(
            user=other_user,
            original_image=build_test_image_base64(color="blue"),
            status="failed",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=130,
        )

        self.client.force_authenticate(user=own_user)
        response = self.client.get(f"/api/statistics/?user_id={other_user.id}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["summary"]["total_records"], 1)
        self.assertEqual(response.data["summary"]["failed_records"], 0)
        self.assertEqual(response.data["filters"]["selected"]["user_id"], own_user.id)

    def test_admin_statistics_api_can_filter_specific_user(self):
        target_user = User.objects.create_user(username="target01")
        other_user = User.objects.create_user(username="other02")
        DetectionRecord.objects.create(
            user=target_user,
            original_image=build_test_image_base64(),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=88,
        )
        DetectionRecord.objects.create(
            user=other_user,
            original_image=build_test_image_base64(color="yellow"),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=95,
        )

        self.client.force_authenticate(user=self.admin)
        response = self.client.get(f"/api/statistics/?user_id={target_user.id}")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["summary"]["total_records"], 1)
        self.assertEqual(response.data["filters"]["selected"]["user_id"], target_user.id)
        self.assertEqual(response.data["charts"]["user_distribution"][0]["label"], "target01")

    def test_miniapp_statistics_can_skip_recent_records(self):
        DetectionRecord.objects.create(
            user=self.admin,
            original_image=build_test_image_base64(),
            original_filename="large-original-name.jpg",
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=88,
        )

        self.client.force_authenticate(user=self.admin)
        response = self.client.get("/api/statistics/?include_recent=0")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["recent_records"], [])
        self.assertIn("summary", response.data)
        self.assertIn("charts", response.data)

    def test_statistics_fast_mode_does_not_read_image_columns(self):
        DetectionRecord.objects.create(
            user=self.admin,
            original_image=build_test_image_base64(),
            visualized_image=build_test_image_base64(color="red"),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=88,
        )

        self.client.force_authenticate(user=self.admin)
        with CaptureQueriesContext(connection) as queries:
            response = self.client.get("/api/statistics/?include_recent=0")

        self.assertEqual(response.status_code, 200)
        self.assertFalse(
            any(
                "original_image" in query["sql"] or "visualized_image" in query["sql"]
                for query in queries
            )
        )

    def test_admin_can_export_statistics_report(self):
        user = User.objects.create_user(username="farmer02")
        record = DetectionRecord.objects.create(
            user=user,
            original_image=build_test_image_base64(),
            status="success",
            model_name="rice-pest-detector",
            model_version="v1",
            duration_ms=110,
        )
        DetectionResult.objects.create(
            record=record,
            label="leaf_blast",
            confidence=0.88,
            xmin=1,
            ymin=1,
            xmax=9,
            ymax=9,
        )

        self.client.force_authenticate(user=self.admin)
        response = self.client.get("/api/admin/statistics/export/?category=leaf_blast")

        self.assertEqual(response.status_code, 200)
        self.assertIn("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        workbook = load_workbook(BytesIO(response.content))
        self.assertIn("概览", workbook.sheetnames)
        self.assertIn("检测明细", workbook.sheetnames)
        self.assertEqual(workbook["概览"]["A1"].value, "检测统计报表")
        detail_values = [cell.value for cell in workbook["检测明细"][3]]
        self.assertIn("leaf_blast", detail_values[-1])
